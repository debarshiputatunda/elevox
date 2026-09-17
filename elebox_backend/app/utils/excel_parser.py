from __future__ import annotations

import re
from dataclasses import dataclass
from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = re.sub(r"[\s_\-]+", " ", text)
    return text


def _cell_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        stripped = value.strip()
        return stripped if stripped else None
    return value


@dataclass(frozen=True)
class ParsedSheet:
    headers: list[str]
    rows: list[dict[str, Any]]
    row_numbers: list[int]


def parse_excel_sheet(
    file_bytes: bytes,
    *,
    required_headers: dict[str, tuple[str, ...]],
    optional_headers: dict[str, tuple[str, ...]] | None = None,
) -> ParsedSheet:
    """
    Parse the first worksheet and map columns using normalized header aliases.

    required_headers maps canonical field names to acceptable header aliases.
    optional_headers maps fields that may be omitted from the spreadsheet.
    """
    optional_headers = optional_headers or {}
    all_headers = {**required_headers, **optional_headers}

    workbook = load_workbook(filename=BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        sheet: Worksheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            raise ValueError("Excel file is empty")

        normalized_headers = [_normalize_header(cell) for cell in header_row]
        alias_to_field: dict[str, str] = {}
        for field, aliases in all_headers.items():
            for alias in aliases:
                alias_to_field[_normalize_header(alias)] = field

        column_map: dict[int, str] = {}
        for index, header in enumerate(normalized_headers):
            if header in alias_to_field:
                column_map[index] = alias_to_field[header]

        missing = [
            field
            for field in required_headers
            if field not in column_map.values()
        ]
        if missing:
            readable = ", ".join(
                required_headers[field][0] for field in missing
            )
            raise ValueError(f"Missing required columns: {readable}")

        parsed_rows: list[dict[str, Any]] = []
        row_numbers: list[int] = []
        excel_row = 1
        for raw_row in rows_iter:
            excel_row += 1
            if raw_row is None or all(
                cell is None or str(cell).strip() == "" for cell in raw_row
            ):
                continue

            record: dict[str, Any] = {}
            for col_index, field in column_map.items():
                record[field] = _cell_value(
                    raw_row[col_index] if col_index < len(raw_row) else None
                )
            parsed_rows.append(record)
            row_numbers.append(excel_row)
    finally:
        workbook.close()

    return ParsedSheet(
        headers=list(all_headers.keys()),
        rows=parsed_rows,
        row_numbers=row_numbers,
    )


def build_template_workbook(headers: list[str]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Import"
    sheet.append(headers)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
